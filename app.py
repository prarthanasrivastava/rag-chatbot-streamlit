from dotenv import load_dotenv
import os
import streamlit as st
import speech_recognition as sr
from transformers import AutoTokenizer, AutoModelForSeq2SeqLM
import torch
from langchain.text_splitter import CharacterTextSplitter
from langchain_community.vectorstores import FAISS
from langchain_community.embeddings import HuggingFaceEmbeddings
from PyPDF2 import PdfReader
from docx import Document
from io import StringIO
from fpdf import FPDF

# Load env vars
load_dotenv()

# Streamlit page config
st.set_page_config(page_title="📚 RAG Chatbot (Local FLAN-T5)", layout="wide", page_icon="🤖")

@st.cache_resource(show_spinner=True)
def load_model():
    tokenizer = AutoTokenizer.from_pretrained("google/flan-t5-base")
    model = AutoModelForSeq2SeqLM.from_pretrained("google/flan-t5-base")
    return tokenizer, model

tokenizer, model = load_model()

def extract_text_from_file(uploaded_file):
    text = ""
    if uploaded_file.name.endswith(".pdf"):
        pdf = PdfReader(uploaded_file)
        for page in pdf.pages:
            text += page.extract_text() or ""
    elif uploaded_file.name.endswith(".docx"):
        doc = Document(uploaded_file)
        for para in doc.paragraphs:
            text += para.text + "\n"
    elif uploaded_file.name.endswith(".txt"):
        stringio = StringIO(uploaded_file.getvalue().decode("utf-8"))
        text = stringio.read()
    return text

def recognize_voice():
    r = sr.Recognizer()
    with sr.Microphone() as source:
        st.toast("🎙 Listening... Speak now", icon="🎤")
        audio = r.listen(source, timeout=5, phrase_time_limit=10)
    try:
        text = r.recognize_google(audio)
        st.toast(f"✅ You said: {text}", icon="✅")
        return text
    except sr.UnknownValueError:
        st.toast("❌ Could not understand audio", icon="❗")
    except sr.RequestError as e:
        st.toast(f"❌ Speech Recognition error: {e}", icon="⚠️")
    return ""

def generate_answer(question, context):
    input_text = f"Context: {context}\nQuestion: {question}\nAnswer:"
    inputs = tokenizer(input_text, return_tensors="pt", truncation=True, max_length=1024)
    outputs = model.generate(**inputs, max_new_tokens=200)
    answer = tokenizer.decode(outputs[0], skip_special_tokens=True)
    return answer

def save_chat_as_pdf(history):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, txt="Chat History", ln=True, align='C')
    for role, msg in history:
        pdf.multi_cell(0, 10, txt=f"{role}: {msg}\n")
    pdf_output = "chat_history.pdf"
    pdf.output(pdf_output)
    return pdf_output

def main():
    # Theme toggle
    theme = st.sidebar.radio("🎨 Theme", ["Light", "Dark"], index=0)
    if theme == "Dark":
        st.markdown("""
            <style>
                body {
                    background-color: #121212;
                    color: #EEEEEE;
                }
                .stTextInput > div > div > input {
                    background-color: #222222;
                    color: #EEEEEE;
                    border: none;
                }
                .stButton>button {
                    background-color: #00ADB5;
                    color: white;
                    border-radius: 8px;
                }
                .custom-box {
                    background-color: #393E46;
                    padding: 16px;
                    border-radius: 10px;
                    color: #EEEEEE;
                }
            </style>
        """, unsafe_allow_html=True)
    else:
        st.markdown("""
            <style>
                body {
                    background-color: #f8f9fa;
                    color: #343a40;
                }
                .stTextInput > div > div > input {
                    background-color: #ffffff;
                    color: #343a40;
                    border: 1px solid #ced4da;
                    border-radius: 10px;
                    padding: 8px;
                }
                .stButton>button {
                    background-color: #0d6efd;
                    color: white;
                    border-radius: 10px;
                    padding: 8px 16px;
                }
                .stButton>button:hover {
                    background-color: #084298;
                }
                .custom-box {
                    background-color: #e9ecef;
                    padding: 16px;
                    border-radius: 10px;
                    margin-top: 10px;
                    box-shadow: 0 4px 6px rgba(0,0,0,0.1);
                }
            </style>
        """, unsafe_allow_html=True)

    st.title("🤖 RAG Chatbot (Local FLAN-T5)")
    st.markdown("<h5 style='color:#6c757d;'>Upload documents and ask your questions using text or voice.</h5>", unsafe_allow_html=True)

    with st.expander("📤 Upload Area", expanded=True):
        uploaded_files = st.file_uploader("Upload PDF, DOCX, or TXT files", type=["pdf", "docx", "txt"], accept_multiple_files=True)

    if uploaded_files:
        full_text = ""
        st.sidebar.markdown("### 📊 Document Summary")
        for f in uploaded_files:
            st.sidebar.markdown(f"📁 **{f.name}**")
            extracted_text = extract_text_from_file(f)
            full_text += extracted_text + "\n"
            snippet = extracted_text[:300].replace("\n", " ") + "..."
            st.sidebar.markdown(f"> {snippet}")

        if not full_text.strip():
            st.warning("❌ Uploaded files are empty or unreadable.")
            return

        splitter = CharacterTextSplitter(chunk_size=800, chunk_overlap=200)
        chunks = splitter.split_text(full_text)

        embeddings = HuggingFaceEmbeddings(model_name="sentence-transformers/all-MiniLM-L6-v2")
        st.session_state.vectorstore = FAISS.from_texts(chunks, embedding=embeddings)
        st.success("✅ Files processed and indexed.")

    col1, col2 = st.columns([5,1])
    with col1:
        query = st.text_input("💬 Ask a question:", placeholder="Type your question here...")
    with col2:
        if st.button("🎙 Voice Input"):
            voice_text = recognize_voice()
            if voice_text:
                query = voice_text

    if query:
        if "vectorstore" not in st.session_state or st.session_state.vectorstore is None:
            st.warning("❗ Please upload and process documents first.")
            return

        docs = st.session_state.vectorstore.similarity_search(query, k=3)
        context = " ".join([doc.page_content for doc in docs])

        try:
            answer = generate_answer(query, context)
            st.session_state.chat_history = st.session_state.get("chat_history", [])
            st.session_state.chat_history.append(("User", query))
            st.session_state.chat_history.append(("Bot", answer))

            st.markdown("### 📜 Answer")
            st.markdown(f"<div class='custom-box'>{answer}</div>", unsafe_allow_html=True)

            with st.expander("🔍 Source Snippets"):
                for doc in docs:
                    snippet = doc.page_content[:300].replace("\n", " ") + "..."
                    st.markdown(f"- `{snippet}`")

        except Exception as e:
            st.error(f"❌ Error during local model inference: {e}")

    if st.session_state.get("chat_history"):
        st.markdown("---")
        st.markdown("### 🧵 Chat History")
        for role, msg in reversed(st.session_state.chat_history):
            name_color = "#0d6efd" if role == "User" else "#343a40"
            st.markdown(f"<b style='color:{name_color};'>{role}:</b> {msg}", unsafe_allow_html=True)

        if st.button("💾 Download Chat as PDF"):
            pdf_path = save_chat_as_pdf(st.session_state.chat_history)
            with open(pdf_path, "rb") as f:
                st.download_button(label="Download PDF", data=f, file_name="chat_history.pdf", mime="application/pdf")

if __name__ == "__main__":
    main()